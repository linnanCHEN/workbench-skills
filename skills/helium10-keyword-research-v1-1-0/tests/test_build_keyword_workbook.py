from __future__ import annotations

import csv
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).parents[1] / "scripts" / "build_keyword_workbook.py"
SPEC = importlib.util.spec_from_file_location("build_keyword_workbook", SCRIPT_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


ASINS = [
    "B000000001",
    "B000000002",
    "B000000003",
    "B000000004",
    "B000000005",
]
SIXTH_ASIN = "B000000006"


class BuildKeywordWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def write_csv(
        self,
        asin: str,
        rows: list[dict[str, object]],
        include_sponsored: bool = True,
    ) -> Path:
        path = self.root / f"{asin}.csv"
        fieldnames = ["phrase", "search_volume", "organic_rank"]
        if include_sponsored:
            fieldnames.append("sponsored_rank")
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        return path

    def build_fixture(self) -> tuple[list[str], dict[str, list[Path]]]:
        rows = {
            ASINS[0]: [
                {"phrase": "alpha toy", "search_volume": 1000, "organic_rank": 5, "sponsored_rank": 2},
                {"phrase": "drop organic", "search_volume": 900, "organic_rank": 25, "sponsored_rank": 1},
                {"phrase": "=formula", "search_volume": 700, "organic_rank": 10, "sponsored_rank": ""},
            ],
            ASINS[1]: [
                {"phrase": "Alpha   Toy", "search_volume": 1100, "organic_rank": 30, "sponsored_rank": 7},
                {"phrase": "drop organic", "search_volume": 900, "organic_rank": "", "sponsored_rank": 1},
            ],
            ASINS[2]: [
                {"phrase": "beta toy", "search_volume": 800, "organic_rank": 24, "sponsored_rank": ""},
                {"phrase": "drop organic", "search_volume": 900, "organic_rank": 99, "sponsored_rank": 1},
            ],
            ASINS[3]: [
                {"phrase": "beta toy", "search_volume": 800, "organic_rank": "", "sponsored_rank": 4},
                {"phrase": "drop organic", "search_volume": 900, "organic_rank": 0, "sponsored_rank": 1},
            ],
        }
        fixture_asins = ASINS[:4]
        sources = {asin: [self.write_csv(asin, rows[asin])] for asin in fixture_asins}
        return fixture_asins, sources

    def build_simple_sources(
        self, count: int, include_sponsored: bool = True
    ) -> tuple[list[str], dict[str, list[Path]]]:
        selected_asins = ASINS[:count]
        sources: dict[str, list[Path]] = {}
        for index, asin in enumerate(selected_asins):
            row: dict[str, object] = {
                "phrase": "shared toy",
                "search_volume": 1000,
                "organic_rank": 1 if index == 0 else 40,
            }
            if include_sponsored:
                row["sponsored_rank"] = 3 if index == 0 else ""
            sources[asin] = [self.write_csv(asin, [row], include_sponsored)]
        return selected_asins, sources

    def test_filter_sort_layout_color_and_formula_safety(self) -> None:
        asins, sources = self.build_fixture()
        records, _counts = MODULE.load_records(asins, sources)
        kept = MODULE.kept_records(records, asins)
        output = self.root / "关键词汇总.xlsx"
        MODULE.build_workbook(kept, asins, output)
        MODULE.validate_output(output, asins, len(kept))

        workbook = load_workbook(output, data_only=False)
        sheet = workbook[MODULE.SHEET_NAME]
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(row=i, column=1).value for i in range(2, 5)], ["alpha toy", "beta toy", "=formula"])
        self.assertEqual(sheet["B2"].value, 1100)
        self.assertEqual(sheet["C2"].value, 5)
        self.assertEqual(sheet["E2"].value, 30)
        self.assertEqual(sheet["F2"].value, 7)
        self.assertEqual(sheet["A4"].data_type, "s")
        self.assertEqual(sheet["A4"].value, "=formula")
        fills = [
            sheet.cell(row=2, column=3 + index * 2).fill.fgColor.rgb
            for index in range(len(asins))
        ]
        self.assertEqual(len(set(fills)), len(asins))

    def test_sponsored_rank_alone_does_not_keep_keyword(self) -> None:
        asins, sources = self.build_fixture()
        records, _counts = MODULE.load_records(asins, sources)
        kept = MODULE.kept_records(records, asins)
        self.assertNotIn("drop organic", [record.display_phrase for record in kept])

    def test_search_volume_conflict_is_detectable(self) -> None:
        asins, sources = self.build_fixture()
        records, _counts = MODULE.load_records(asins, sources)
        alpha = records["alpha toy"]
        self.assertEqual(alpha.search_volumes, {1000, 1100})

    def test_missing_column_is_rejected(self) -> None:
        bad_path = self.root / "bad.csv"
        bad_path.write_text(
            "phrase,search_volume,sponsored_rank\nalpha,10,1\n", encoding="utf-8"
        )
        sources = {asin: [bad_path] for asin in ASINS[:2]}
        with self.assertRaisesRegex(ValueError, "missing required columns"):
            MODULE.load_records(ASINS[:2], sources)

    def test_missing_sponsored_column_is_allowed(self) -> None:
        asins, sources = self.build_simple_sources(2, include_sponsored=False)
        records, _counts = MODULE.load_records(asins, sources)
        kept = MODULE.kept_records(records, asins)
        output = self.root / "自然排名汇总.xlsx"
        MODULE.build_workbook(kept, asins, output)
        MODULE.validate_output(output, asins, len(kept))

        sheet = load_workbook(output, data_only=False)[MODULE.SHEET_NAME]
        self.assertEqual(sheet["C2"].value, 1)
        self.assertIsNone(sheet["D2"].value)

    def test_two_to_five_asins_are_supported(self) -> None:
        for count in range(2, 6):
            with self.subTest(count=count):
                asins, sources = self.build_simple_sources(count)
                source_args = [f"{asin}={sources[asin][0]}" for asin in asins]
                parsed_asins, parsed_sources = MODULE.parse_sources(source_args)
                self.assertEqual(parsed_asins, asins)

                records, _counts = MODULE.load_records(parsed_asins, parsed_sources)
                kept = MODULE.kept_records(records, parsed_asins)
                output = self.root / f"关键词汇总-{count}.xlsx"
                MODULE.build_workbook(kept, parsed_asins, output)
                MODULE.validate_output(output, parsed_asins, len(kept))

                sheet = load_workbook(output, data_only=False)[MODULE.SHEET_NAME]
                self.assertEqual(sheet.max_column, 2 + count * 2)
                fills = [
                    sheet.cell(row=2, column=3 + index * 2).fill.fgColor.rgb
                    for index in range(count)
                ]
                self.assertEqual(len(set(fills)), count)
                self.assertTrue(sheet["C1"].font.bold)
                self.assertFalse(sheet["D1"].font.bold)

    def test_asin_count_outside_range_is_rejected(self) -> None:
        one_asin, one_source = self.build_simple_sources(1)
        with self.assertRaisesRegex(ValueError, "Between 2 and 5"):
            MODULE.parse_sources([f"{one_asin[0]}={one_source[one_asin[0]][0]}"])

        five_asins, five_sources = self.build_simple_sources(5)
        sixth_path = self.write_csv(
            SIXTH_ASIN,
            [
                {
                    "phrase": "shared toy",
                    "search_volume": 1000,
                    "organic_rank": 50,
                    "sponsored_rank": "",
                }
            ],
        )
        six_args = [f"{asin}={five_sources[asin][0]}" for asin in five_asins]
        six_args.append(f"{SIXTH_ASIN}={sixth_path}")
        with self.assertRaisesRegex(ValueError, "Between 2 and 5"):
            MODULE.parse_sources(six_args)

    def test_receipt_json_can_preserve_unicode(self) -> None:
        receipt = self.root / "关键词汇总收据.json"
        receipt.write_text(json.dumps({"状态": "通过"}, ensure_ascii=False), encoding="utf-8")
        self.assertEqual(json.loads(receipt.read_text(encoding="utf-8"))["状态"], "通过")

    def test_cli_generates_workbook_and_receipt(self) -> None:
        asins, sources = self.build_fixture()
        output = self.root / "关键词汇总.xlsx"
        receipt = self.root / "关键词汇总收据.json"
        command = [sys.executable, str(SCRIPT_PATH)]
        for asin in asins:
            command.extend(["--source", f"{asin}={sources[asin][0]}"])
        command.extend(["--output", str(output), "--receipt", str(receipt)])

        result = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(receipt.read_text(encoding="utf-8"))
        self.assertEqual(payload["status"], "PASS_WITH_LIMITS")
        self.assertEqual(payload["asin_count"], 4)
        self.assertEqual(payload["primary_metric"], "organic_rank")
        self.assertEqual(payload["secondary_metric"], "sponsored_rank_optional")
        self.assertEqual(payload["kept_keywords"], 3)
        self.assertEqual(payload["deleted_keywords"], 1)
        self.assertTrue(output.is_file())

    def test_cli_supports_two_and_five_asins(self) -> None:
        for count in (2, 5):
            with self.subTest(count=count):
                asins, sources = self.build_simple_sources(count)
                output = self.root / f"关键词汇总-cli-{count}.xlsx"
                receipt = self.root / f"关键词汇总收据-cli-{count}.json"
                command = [sys.executable, str(SCRIPT_PATH)]
                for asin in asins:
                    command.extend(["--source", f"{asin}={sources[asin][0]}"])
                command.extend(
                    ["--output", str(output), "--receipt", str(receipt)]
                )

                result = subprocess.run(
                    command, capture_output=True, text=True, check=False
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                payload = json.loads(receipt.read_text(encoding="utf-8"))
                self.assertEqual(payload["status"], "PASS")
                self.assertEqual(payload["asin_count"], count)
                self.assertTrue(output.is_file())


if __name__ == "__main__":
    unittest.main()

