#!/usr/bin/env python3
"""Build a deterministic two-to-five-ASIN Helium 10 keyword workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$")
MIN_ASINS = 2
MAX_ASINS = 5
REQUIRED_COLUMNS = {"phrase", "search_volume", "organic_rank"}
SHEET_NAME = "关键词汇总"
ORGANIC_THRESHOLD = 24
ASIN_PALETTES = (
    ("1F4E78", "D9EAF7", "EDF5FA"),
    ("548235", "E2F0D9", "F1F7ED"),
    ("A64D00", "FCE4D6", "FDF2EB"),
    ("7030A0", "E4DFEC", "F2EFF6"),
    ("0F6B78", "DDEBF7", "EEF5F8"),
)


@dataclass
class KeywordRecord:
    display_phrase: str
    search_volumes: set[int] = field(default_factory=set)
    organic_ranks: dict[str, int | None] = field(default_factory=dict)
    sponsored_ranks: dict[str, int | None] = field(default_factory=dict)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge 2-5 Helium 10 CSV exports into a natural-rank-focused Excel workbook."
    )
    parser.add_argument(
        "--source",
        action="append",
        required=True,
        metavar="ASIN=CSV",
        help="Map an ASIN to one CSV page. Repeat for more pages.",
    )
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx path.")
    parser.add_argument(
        "--receipt", required=True, type=Path, help="Output JSON receipt path."
    )
    return parser.parse_args()


def normalize_phrase(value: str) -> tuple[str, str]:
    display = " ".join(unicodedata.normalize("NFKC", value).strip().split())
    return display.casefold(), display


def parse_positive_int(value: object) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text.replace(",", ""))
    except ValueError:
        return None
    if not math.isfinite(number) or number <= 0:
        return None
    return int(number)


def best_rank(existing: int | None, candidate: int | None) -> int | None:
    if existing is None:
        return candidate
    if candidate is None:
        return existing
    return min(existing, candidate)


def parse_sources(values: Iterable[str]) -> tuple[list[str], dict[str, list[Path]]]:
    sources: dict[str, list[Path]] = OrderedDict()
    for raw in values:
        if "=" not in raw:
            raise ValueError(f"Invalid --source value: {raw!r}; expected ASIN=CSV")
        asin_raw, path_raw = raw.split("=", 1)
        asin = asin_raw.strip().upper()
        if not ASIN_PATTERN.fullmatch(asin):
            raise ValueError(f"Invalid ASIN: {asin_raw!r}")
        path = Path(path_raw).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"CSV source not found: {path}")
        sources.setdefault(asin, []).append(path)
    if not MIN_ASINS <= len(sources) <= MAX_ASINS:
        raise ValueError(
            f"Between {MIN_ASINS} and {MAX_ASINS} unique ASINs are required; got {len(sources)}"
        )
    return list(sources), sources


def load_records(
    asins: list[str], sources: dict[str, list[Path]]
) -> tuple[dict[str, KeywordRecord], dict[str, int]]:
    records: dict[str, KeywordRecord] = OrderedDict()
    row_counts: dict[str, int] = defaultdict(int)

    for asin in asins:
        for source in sources[asin]:
            with source.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                fields = {field.strip() for field in (reader.fieldnames or []) if field}
                missing = REQUIRED_COLUMNS - fields
                if missing:
                    raise ValueError(
                        f"{source} is missing required columns: {', '.join(sorted(missing))}"
                    )
                for row in reader:
                    row_counts[asin] += 1
                    key, display = normalize_phrase(row.get("phrase", ""))
                    if not key:
                        continue
                    record = records.setdefault(key, KeywordRecord(display_phrase=display))
                    volume = parse_positive_int(row.get("search_volume"))
                    if volume is not None:
                        record.search_volumes.add(volume)
                    organic = parse_positive_int(row.get("organic_rank"))
                    sponsored = parse_positive_int(row.get("sponsored_rank"))
                    record.organic_ranks[asin] = best_rank(
                        record.organic_ranks.get(asin), organic
                    )
                    record.sponsored_ranks[asin] = best_rank(
                        record.sponsored_ranks.get(asin), sponsored
                    )

    return records, dict(row_counts)


def kept_records(records: dict[str, KeywordRecord], asins: list[str]) -> list[KeywordRecord]:
    kept = [
        record
        for record in records.values()
        if any(
            rank is not None and 1 <= rank <= ORGANIC_THRESHOLD
            for rank in (record.organic_ranks.get(asin) for asin in asins)
        )
    ]
    return sorted(
        kept,
        key=lambda record: (
            -(max(record.search_volumes) if record.search_volumes else -1),
            record.display_phrase.casefold(),
        ),
    )


def write_text_cell(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def build_workbook(records: list[KeywordRecord], asins: list[str], output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    headers = ["关键词", "搜索量"]
    for asin in asins:
        headers.extend([f"{asin} 自然排名", f"{asin} 广告排名"])
    sheet.append(headers)

    neutral_header = PatternFill("solid", fgColor="404040")
    white_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E1F2")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        cell.fill = neutral_header

    for index, asin in enumerate(asins):
        dark, _organic_light, _sponsored_light = ASIN_PALETTES[index]
        organic_column = 3 + index * 2
        sponsored_column = organic_column + 1
        sheet.cell(row=1, column=organic_column).fill = PatternFill("solid", fgColor=dark)
        sheet.cell(row=1, column=organic_column).comment = Comment(
            f"{asin} 自然排名；这是保留关键词的核心过滤字段，空白表示未检测到排名。",
            "Codex",
        )
        sheet.cell(row=1, column=sponsored_column).fill = PatternFill("solid", fgColor=dark)
        sheet.cell(row=1, column=sponsored_column).font = Font(color="FFFFFF", bold=False)
        sheet.cell(row=1, column=sponsored_column).comment = Comment(
            f"{asin} 广告排名；仅作辅助，字段或值缺失时保持空白。", "Codex"
        )

    for row_index, record in enumerate(records, start=2):
        write_text_cell(sheet.cell(row=row_index, column=1), record.display_phrase)
        volume = max(record.search_volumes) if record.search_volumes else None
        sheet.cell(row=row_index, column=2, value=volume)
        if len(record.search_volumes) > 1:
            observed = ", ".join(str(value) for value in sorted(record.search_volumes))
            sheet.cell(row=row_index, column=2).comment = Comment(
                f"多个 ASIN 导出出现不同搜索量：{observed}。本表取最大值用于排序；请人工复核。",
                "Codex",
            )

        for index, asin in enumerate(asins):
            organic_column = 3 + index * 2
            sponsored_column = organic_column + 1
            sheet.cell(
                row=row_index,
                column=organic_column,
                value=record.organic_ranks.get(asin),
            )
            sheet.cell(
                row=row_index,
                column=sponsored_column,
                value=record.sponsored_ranks.get(asin),
            )
            _dark, organic_light, sponsored_light = ASIN_PALETTES[index]
            sheet.cell(row=row_index, column=organic_column).fill = PatternFill(
                "solid", fgColor=organic_light
            )
            sheet.cell(row=row_index, column=organic_column).font = Font(bold=True)
            sheet.cell(row=row_index, column=sponsored_column).fill = PatternFill(
                "solid", fgColor=sponsored_light
            )

        for column in range(1, len(headers) + 1):
            cell = sheet.cell(row=row_index, column=column)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "center", vertical="center"
            )

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    sheet.row_dimensions[1].height = 28
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 14
    for column in range(3, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 19

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_output(output: Path, asins: list[str], expected_rows: int) -> None:
    workbook = load_workbook(output, data_only=False)
    if workbook.sheetnames != [SHEET_NAME]:
        raise RuntimeError(f"Unexpected worksheets: {workbook.sheetnames}")
    sheet = workbook[SHEET_NAME]
    if sheet.max_row != expected_rows + 1:
        raise RuntimeError(
            f"Workbook row count mismatch: expected {expected_rows + 1}, got {sheet.max_row}"
        )
    for row in range(2, sheet.max_row + 1):
        organic_ranks = [
            sheet.cell(row=row, column=3 + index * 2).value
            for index in range(len(asins))
        ]
        if not any(isinstance(rank, int) and 1 <= rank <= ORGANIC_THRESHOLD for rank in organic_ranks):
            raise RuntimeError(f"Row {row} violates the organic-rank filter")
    expected_headers = ["关键词", "搜索量"]
    for asin in asins:
        expected_headers.extend([f"{asin} 自然排名", f"{asin} 广告排名"])
    actual_headers = [
        sheet.cell(row=1, column=column).value
        for column in range(1, len(expected_headers) + 1)
    ]
    if actual_headers != expected_headers:
        raise RuntimeError("Workbook headers do not match the required layout")


def main() -> int:
    args = parse_args()
    try:
        asins, sources = parse_sources(args.source)
        records, row_counts = load_records(asins, sources)
        kept = kept_records(records, asins)
        build_workbook(kept, asins, args.output)
        validate_output(args.output, asins, len(kept))

        conflicts = [
            {
                "keyword": record.display_phrase,
                "observed_search_volumes": sorted(record.search_volumes),
                "selected_search_volume": max(record.search_volumes),
            }
            for record in kept
            if len(record.search_volumes) > 1
        ]
        receipt = {
            "status": "PASS_WITH_LIMITS" if conflicts else "PASS",
            "asins": asins,
            "asin_count": len(asins),
            "source_files": {
                asin: [str(path) for path in paths] for asin, paths in sources.items()
            },
            "source_row_counts": row_counts,
            "unique_keywords_before_filter": len(records),
            "kept_keywords": len(kept),
            "deleted_keywords": len(records) - len(kept),
            "organic_rank_threshold": ORGANIC_THRESHOLD,
            "primary_metric": "organic_rank",
            "secondary_metric": "sponsored_rank_optional",
            "search_volume_conflict_policy": "max_non_null_with_cell_comment",
            "search_volume_conflicts": conflicts,
            "workbook_path": str(args.output.resolve()),
            "workbook_sha256": sha256_file(args.output),
        }
        args.receipt.parent.mkdir(parents=True, exist_ok=True)
        args.receipt.write_text(
            json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(receipt, ensure_ascii=False))
        return 0
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

